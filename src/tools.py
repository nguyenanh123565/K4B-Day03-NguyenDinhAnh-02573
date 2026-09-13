"""
🛠️ TOOL DEFINITIONS & EXECUTION BACKEND
Mã nguồn chứa danh sách Tool Schemas (JSON Schema) vàa Execution Layer phục vụ cho MCP Server.
"""

import json
import os
from datetime import datetime
from typing import Dict, Any
from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook

load_dotenv()

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
WORKBOOK_PATH = os.path.join(BASE_DIR, os.getenv("EXPENSE_WORKBOOK", "data/expense_tracker.xlsx"))
TRANSACTION_HEADERS = [
    "id", "date", "type", "amount", "category", "subcategory", "note",
    "payment_method", "account", "tags", "is_recurring", "created_at"
]

# ==============================================================================
# 1. KHAI BÁO TOOL SCHEMAS CHUẨN NATIVE JSON SCHEMA (TASK 1.2)
# ==============================================================================

TOOLS_SCHEMA = [
    {
        "name": "record_expense",
        "description": "Ghi nhận một khoản chi tiêu cá nhân vào sổ chi tiêu.",
        "parameters": {
            "type": "object",
            "properties": {
                "amount": {"type": "number", "description": "Số tiền đã chi bằng VND."},
                "category": {"type": "string", "description": "Danh mục chi tiêu, ví dụ Ăn uống hoặc Di chuyển."},
                "description": {"type": "string", "description": "Mô tả ngắn khoản chi."},
                "date": {"type": "string", "description": "Ngày chi theo định dạng YYYY-MM-DD."}
            },
            "required": ["amount", "category", "description", "date"]
        }
    },
    {
        "name": "summarize_expenses",
        "description": "Tổng hợp các khoản chi theo tháng hoặc danh mục.",
        "parameters": {
            "type": "object",
            "properties": {
                "month": {"type": "string", "description": "Tháng cần tổng hợp theo định dạng YYYY-MM."},
                "category": {"type": "string", "description": "Danh mục cần lọc; bỏ trống để xem tất cả."}
            },
            "required": ["month"]
        }
    },
    {
        "name": "budget_check",
        "description": "Kiểm tra số tiền đã chi so với ngân sách tháng của một danh mục.",
        "parameters": {
            "type": "object",
            "properties": {
                "category": {"type": "string", "description": "Danh mục cần kiểm tra."},
                "monthly_limit": {"type": "number", "description": "Ngân sách tối đa trong tháng bằng VND."},
                "month": {"type": "string", "description": "Tháng cần kiểm tra theo định dạng YYYY-MM."}
            },
            "required": ["category", "monthly_limit", "month"]
        }
    }
]

# ==============================================================================
# 2. MÔ PHỎNG DỮ LIỆU & HÀM THỰC THI TOOL (EXECUTION LAYER)
# ==============================================================================

def _ensure_workbook() -> None:
    """Create the documented workbook structure when no workbook exists."""
    if os.path.exists(WORKBOOK_PATH):
        return
    os.makedirs(os.path.dirname(WORKBOOK_PATH), exist_ok=True)
    workbook = Workbook()
    transactions = workbook.active
    transactions.title = "Transactions"
    transactions.append(TRANSACTION_HEADERS)
    transactions.append([1, "2026-09-10", "expense", 85000, "Ăn uống", "", "Cơm trưa", "Tiền mặt", "", "", False, datetime.now().isoformat(timespec="seconds")])
    transactions.append([2, "2026-09-11", "expense", 40000, "Di chuyển", "", "Xe bus", "Tiền mặt", "", "", False, datetime.now().isoformat(timespec="seconds")])
    categories = workbook.create_sheet("Categories")
    categories.append(["category_name", "type", "parent_category", "icon_color"])
    budget = workbook.create_sheet("Budget")
    budget.append(["category", "period", "limit_amount", "start_date", "note"])
    budget.append(["Ăn uống", "monthly", 3000000, "2026-09-01", "Ngân sách mẫu"])
    workbook.create_sheet("Accounts").append(["account_name", "account_type", "initial_balance", "currency"])
    workbook.create_sheet("Report").append(["period", "category", "total_amount", "percent_of_total", "budget_status", "generated_at"])
    workbook.save(WORKBOOK_PATH)


def _read_transactions() -> list:
    _ensure_workbook()
    workbook = load_workbook(WORKBOOK_PATH, data_only=True)
    sheet = workbook["Transactions"]
    headers = [cell.value for cell in sheet[1]]
    return [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True) if any(row)]


def execute_record_expense(amount: float, category: str, description: str, date: str) -> str:
    """Ghi nhận khoản chi sau khi kiểm tra dữ liệu đầu vào."""
    if amount <= 0:
        return json.dumps({"status": "VALIDATION_ERROR", "message": "Số tiền phải lớn hơn 0."}, ensure_ascii=False)
    datetime.strptime(date, "%Y-%m-%d")
    _ensure_workbook()
    workbook = load_workbook(WORKBOOK_PATH)
    sheet = workbook["Transactions"]
    next_id = max((row[0].value or 0 for row in sheet.iter_rows(min_row=2, max_col=1)), default=0) + 1
    expense = {"id": next_id, "date": date, "type": "expense", "amount": amount,
               "category": category.strip(), "note": description.strip()}
    sheet.append([next_id, date, "expense", amount, category.strip(), "", description.strip(), "", "", "", False,
                  datetime.now().isoformat(timespec="seconds")])
    workbook.save(WORKBOOK_PATH)
    return json.dumps({"status": "SUCCESS", "event": "EXPENSE_RECORDED", "data": expense,
                       "message": f"Đã ghi nhận khoản chi {amount:,.0f} VND cho danh mục {category}."}, ensure_ascii=False)


def execute_summarize_expenses(month: str, category: str = "") -> str:
    """Tổng hợp khoản chi theo tháng và tùy chọn danh mục."""
    expenses = [e for e in _read_transactions() if str(e.get("date", "")).startswith(month)
                and e.get("type") == "expense"
                and (not category or str(e.get("category", "")).lower() == category.strip().lower())]
    expenses = [{"id": e.get("id"), "amount": e.get("amount"), "category": e.get("category"),
                 "description": e.get("note", ""), "date": str(e.get("date", ""))} for e in expenses]
    if not expenses:
        return json.dumps({"status": "NOT_FOUND", "message": f"Chưa có giao dịch nào trong tháng {month} cho danh mục '{category}'."}, ensure_ascii=False)
    total = sum(e["amount"] for e in expenses)
    return json.dumps({"status": "SUCCESS", "month": month, "category": category or "Tất cả",
                       "total": total, "count": len(expenses), "expenses": expenses}, ensure_ascii=False)


def execute_check_budget(category: str, monthly_limit: float, month: str) -> str:
    """So sánh tổng chi của danh mục với ngân sách tháng."""
    _ensure_workbook()
    workbook = load_workbook(WORKBOOK_PATH, data_only=True)
    budget_sheet = workbook["Budget"]
    for row in budget_sheet.iter_rows(min_row=2, values_only=True):
        if row[0] == category and row[1] == "monthly" and row[2] is not None:
            monthly_limit = float(row[2])
            break
    summary = json.loads(execute_summarize_expenses(month, category))
    spent = summary.get("total", 0)
    return json.dumps({"status": "SUCCESS", "month": month, "category": category,
                       "spent": spent, "monthly_limit": monthly_limit,
                       "remaining": monthly_limit - spent,
                       "over_budget": spent > monthly_limit}, ensure_ascii=False)


# Router gọi tool thực tế
TOOL_ROUTER = {
    "record_expense": execute_record_expense,
    "summarize_expenses": execute_summarize_expenses,
    "budget_check": execute_check_budget
}

def dispatch_tool_call(tool_name: str, arguments: Dict[str, Any]) -> str:
    """Hàm trung chuyển thực thi tool"""
    if tool_name in TOOL_ROUTER:
        try:
            return TOOL_ROUTER[tool_name](**arguments)
        except Exception as e:
            return json.dumps({"status": "EXECUTION_ERROR", "error": str(e)}, ensure_ascii=False)
    return json.dumps({"status": "UNKNOWN_TOOL", "error": f"Tool '{tool_name}' không tồn tại!"}, ensure_ascii=False)
